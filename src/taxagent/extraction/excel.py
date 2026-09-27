"""Read the payroll register cell by cell. Never OCR Excel: the cells are exact."""

from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from taxagent.schema.models import ExtractedValue, SourceRef

CENT = Decimal("0.01")

# Register column header -> canonical field name
MONEY_COLUMNS = {
    "Regular Wages": "register.regular_wages",
    "Bonus": "register.bonus",
    "Gross Wages": "register.gross_wages",
    "Federal Tax": "register.federal_tax",
    "Social Security Tax": "register.ss_tax",
    "Medicare Tax": "register.medicare_tax",
}


def to_money(raw: object) -> tuple[Decimal | None, float]:
    """Turn a cell's content into (amount, confidence)."""
    if isinstance(raw, int | float):
        # Via str(): Decimal(458.56) would keep the float's binary error
        return Decimal(str(raw)).quantize(CENT), 1.0
    if isinstance(raw, str):
        try:
            return Decimal(raw.replace(",", "")).quantize(CENT), 0.9  # typed as text
        except InvalidOperation:
            pass
    return None, 0.0


def extract_register(path: Path, file_sha256: str) -> list[ExtractedValue]:
    sheet = load_workbook(path, data_only=True).active

    # Real exports have title rows, so find the header instead of assuming row 1
    header_row = next(row for row in sheet.iter_rows() if row[0].value == "Employee ID")
    columns = {cell.column: cell.value for cell in header_row}

    values = []
    for row in sheet.iter_rows(min_row=header_row[0].row + 1):
        employee_id = row[0].value
        if employee_id in (None, "TOTAL"):
            continue
        for cell in row:
            field = MONEY_COLUMNS.get(columns.get(cell.column))
            if field is None:
                continue  # name, SSN: not numbers we reconcile
            amount, confidence = to_money(cell.value)
            values.append(
                ExtractedValue(
                    field=field,
                    employee_id=employee_id,
                    value=amount,
                    confidence=confidence,
                    source=SourceRef(
                        file_sha256=file_sha256, sheet=sheet.title, cell=cell.coordinate
                    ),
                )
            )
    return values
