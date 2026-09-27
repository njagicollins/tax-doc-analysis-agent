from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from synthetic.employees import Employee, generate_employees
from synthetic.payroll import QUARTERS, PayrollLine, generate_payroll

OUTPUT_DIR = Path("data/sources/adp")

# Planted error #2: one employee is left out of the Q1 register
MISSING_EMPLOYEE_ID = "E013"
MISSING_QUARTER = "2026Q1"

HEADERS = [
    "Employee ID",
    "Name",
    "SSN",
    "Regular Wages",
    "Bonus",
    "Gross Wages",
    "Federal Tax",
    "Social Security Tax",
    "Medicare Tax",
]
MONEY_FIELDS = [
    "regular_wages",
    "bonus",
    "gross_wages",
    "federal_tax",
    "ss_tax",
    "medicare_tax",
]
MONEY_FORMAT = "#,##0.00"


def write_register(
    quarter: str,
    employees: list[Employee],
    lines: list[PayrollLine],
    output_dir: Path = OUTPUT_DIR,
) -> Path:
    by_id = {e.employee_id: e for e in employees}

    rows = [
        line
        for line in lines
        if line.quarter == quarter
        and not (line.employee_id == MISSING_EMPLOYEE_ID and quarter == MISSING_QUARTER)
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Register"

    # Real exports rarely start with headers on row 1
    ws.append([f"Payroll Register - {quarter}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(HEADERS)
    for cell in ws[3]:
        cell.font = Font(bold=True)

    for line in rows:
        e = by_id[line.employee_id]
        ws.append(
            [e.employee_id, e.name, e.ssn] + [getattr(line, f) for f in MONEY_FIELDS]
        )

    totals = [sum(getattr(line, f) for line in rows) for f in MONEY_FIELDS]
    ws.append(["TOTAL", "", ""] + totals)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    for row in ws.iter_rows(min_row=4, min_col=4, max_col=9):
        for cell in row:
            cell.number_format = MONEY_FORMAT

    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"payroll_register_{quarter}.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    employees = generate_employees()
    lines = generate_payroll(employees)
    for quarter in QUARTERS:
        print(write_register(quarter, employees, lines))
